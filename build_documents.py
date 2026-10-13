import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Ordered list of screenshots corresponding to the 27 placeholders
SCREENSHOT_FILES = [
    # Practical 1
    ("p1_ss1_env_setup.png", "Project folder, activated virtual environment and successful dependency installation."),
    ("p1_ss2_train_model.png", "Training output, evaluation metrics and model.pkl created successfully."),
    ("p1_ss3_flask_run.png", "Running Flask web interface showing server output and endpoints."),
    ("p1_ss4_prediction_web.png", "Prediction result handled via web interface endpoint."),
    ("p1_ss5_health_api.png", "Successful health check and prediction API JSON response."),
    # Practical 2
    ("p2_ss1_docker_build.png", "Successful Docker image build output."),
    ("p2_ss2_docker_images.png", "docker images output verifying student-marks:1.0 image."),
    ("p2_ss3_docker_ps.png", "docker ps and docker logs output showing the running container."),
    ("p2_ss4_docker_access.png", "Flask application accessed through the Docker container via curl."),
    # Practical 3
    ("p3_ss1_ec2_launch.png", "EC2 launch configuration showing instance name, Ubuntu AMI, instance type and key pair."),
    ("p3_ss2_ec2_sg.png", "EC2 security group inbound rules allowing ports 22 and 5000."),
    ("p3_ss3_ssh_connect.png", "Successful SSH connection to the Ubuntu instance."),
    ("p3_ss4_docker_install.png", "Docker installed and active on AWS EC2."),
    ("p3_ss5_transfer_files.png", "Project files transferred and loaded successfully on EC2 via SCP."),
    ("p3_ss6_docker_run_ec2.png", "Running application container and images on EC2."),
    ("p3_ss7_verify_ec2.png", "Deployed application verified through EC2 public IP."),
    ("p3_ss8_predict_ec2.png", "Successful prediction through the AWS EC2 deployment."),
    # Practical 4
    ("p4_ss1_gcp_project.png", "Newly created GCP project verified in Google Cloud Shell."),
    ("p4_ss2_gcp_compute.png", "Compute Engine API enabled with active project configuration."),
    ("p4_ss3_gcp_vm.png", "Compute Engine VM configuration showing name, zone, machine type and external IP."),
    ("p4_ss4_gcp_firewall.png", "GCP firewall rule showing TCP port 5000 and configured 0.0.0.0/0 source range."),
    ("p4_ss5_gcp_ssh.png", "SSH terminal connected to the GCP VM instance."),
    ("p4_ss6_gcp_docker_install.png", "Docker installation and daemon verified on GCP VM."),
    ("p4_ss7_gcp_docker_import.png", "Project files and Docker environment imported on GCP VM."),
    ("p4_ss8_gcp_docker_run.png", "Running application container on GCP Compute Engine VM."),
    ("p4_ss9_gcp_verify.png", "ML application opened and verified using GCP external IP."),
    ("p4_ss10_gcp_predict.png", "Successful prediction response through GCP deployment."),
]

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def create_docx_report(prompt_file, screenshots_dir, output_docx):
    with open(prompt_file, 'r', encoding='utf-8') as f:
        text = f.read()

    # Clean header tags
    text = re.sub(r'<\/?USER_REQUEST>', '', text)
    text = re.sub(r'<ADDITIONAL_METADATA>[\s\S]*?<\/ADDITIONAL_METADATA>', '', text)
    text = re.sub(r'DO THESE STEPS AND GENERATE SCREENSHOTS[\s\S]*', '', text).strip()

    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("MLOps LAB PRACTICAL MANUAL & REPORT")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("Practicals 1 to 4: Model Serialization, Dockerization, AWS EC2 & GCP Deployment\nComplete Command-Line Verification & Terminal Outputs")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(11)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    # Split text by placeholder
    pattern = r'\[SCREENSHOT PLACEHOLDER\s*\]'
    segments = re.split(pattern, text)

    current_ss_idx = 0

    for idx, segment in enumerate(segments):
        # Process lines in segment
        lines = segment.split('\n')
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line:
                i += 1
                continue

            if re.match(r'^PRACTICAL\s+\d', line, re.IGNORECASE):
                doc.add_page_break() if idx > 0 or i > 5 else None
                h1 = doc.add_paragraph()
                h1.paragraph_format.space_before = Pt(16)
                h1.paragraph_format.space_after = Pt(6)
                r = h1.add_run(line)
                r.font.name = "Calibri"
                r.font.size = Pt(18)
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 58, 138)
            elif line.startswith("AIM:") or line.startswith("AIM :"):
                aim_p = doc.add_paragraph()
                aim_p.paragraph_format.space_before = Pt(6)
                aim_p.paragraph_format.space_after = Pt(6)
                r_label = aim_p.add_run("AIM: ")
                r_label.font.bold = True
                r_label.font.color.rgb = RGBColor(15, 23, 42)
                r_val = aim_p.add_run(line[4:].strip())
                r_val.font.color.rgb = RGBColor(51, 65, 85)
            elif re.match(r'^Step\s+\d', line, re.IGNORECASE):
                h2 = doc.add_paragraph()
                h2.paragraph_format.space_before = Pt(12)
                h2.paragraph_format.space_after = Pt(4)
                r = h2.add_run(line)
                r.font.name = "Calibri"
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(14, 116, 144)  # Teal/Blue
            elif any(line.startswith(c) for c in ["python ", "pip ", "docker ", "sudo ", "curl ", "gcloud ", "aws ", "ssh ", "scp ", "mkdir ", "cd "]):
                # Code block style
                code_table = doc.add_table(rows=1, cols=1)
                code_table.autofit = False
                code_table.columns[0].width = Inches(6.8)
                cell = code_table.cell(0, 0)
                set_cell_background(cell, "F1F5F9")
                cell_p = cell.paragraphs[0]
                cell_p.paragraph_format.space_before = Pt(2)
                cell_p.paragraph_format.space_after = Pt(2)
                cr = cell_p.add_run(line)
                cr.font.name = "Consolas"
                cr.font.size = Pt(9.5)
                cr.font.color.rgb = RGBColor(15, 23, 42)
            else:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                r = p.add_run(line)
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(51, 65, 85)
            i += 1

        # Now insert screenshot if not at the very last segment
        if idx < len(segments) - 1 and current_ss_idx < len(SCREENSHOT_FILES):
            ss_filename, ss_caption = SCREENSHOT_FILES[current_ss_idx]
            ss_path = os.path.join(screenshots_dir, ss_filename)
            if os.path.exists(ss_path):
                img_p = doc.add_paragraph()
                img_p.paragraph_format.space_before = Pt(8)
                img_p.paragraph_format.space_after = Pt(3)
                img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                img_p.add_run().add_picture(ss_path, width=Inches(6.6))

                cap_p = doc.add_paragraph()
                cap_p.paragraph_format.space_before = Pt(0)
                cap_p.paragraph_format.space_after = Pt(12)
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_run = cap_p.add_run(f"Figure {current_ss_idx + 1}: {ss_caption}")
                cap_run.font.name = "Calibri"
                cap_run.font.size = Pt(9)
                cap_run.font.bold = True
                cap_run.font.italic = True
                cap_run.font.color.rgb = RGBColor(71, 85, 105)

            current_ss_idx += 1

    doc.save(output_docx)
    print(f"Successfully generated DOCX manual: {output_docx}")

def create_html_report(prompt_file, screenshots_dir, output_html):
    with open(prompt_file, 'r', encoding='utf-8') as f:
        text = f.read()

    text = re.sub(r'<\/?USER_REQUEST>', '', text)
    text = re.sub(r'<ADDITIONAL_METADATA>[\s\S]*?<\/ADDITIONAL_METADATA>', '', text)
    text = re.sub(r'DO THESE STEPS AND GENERATE SCREENSHOTS[\s\S]*', '', text).strip()

    pattern = r'\[SCREENSHOT PLACEHOLDER\s*\]'
    segments = re.split(pattern, text)

    html_parts = []
    html_parts.append("""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>MLOps Practicals 1-4 Complete Lab Manual & Terminal Outputs</title>
  <style>
    body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; background: #0f172a; color: #f1f5f9; line-height: 1.6; margin: 0; padding: 2rem; }
    .container { max-width: 960px; margin: 0 auto; background: #1e293b; padding: 2.5rem 3rem; border-radius: 12px; box-shadow: 0 10px 30px rgba(0,0,0,0.5); border: 1px solid #334155; }
    h1.doc-title { font-size: 2.2rem; color: #38bdf8; text-align: center; margin-bottom: 0.2rem; border-bottom: 2px solid #38bdf8; padding-bottom: 0.8rem; }
    p.subtitle { text-align: center; color: #94a3b8; font-style: italic; margin-bottom: 2.5rem; }
    h2.practical-title { font-size: 1.6rem; color: #60a5fa; border-left: 4px solid #60a5fa; padding-left: 12px; margin-top: 2.5rem; margin-bottom: 1rem; }
    h3.step-title { font-size: 1.2rem; color: #34d399; margin-top: 1.8rem; margin-bottom: 0.5rem; }
    .aim-box { background: rgba(56, 189, 248, 0.1); border-left: 4px solid #38bdf8; padding: 10px 16px; border-radius: 4px; margin: 12px 0; }
    .code-box { background: #090d16; color: #f8fafc; font-family: Consolas, monospace; font-size: 0.95rem; padding: 10px 14px; border-radius: 6px; border: 1px solid #1e293b; margin: 8px 0; overflow-x: auto; }
    .figure-box { text-align: center; margin: 1.5rem 0; background: #0b0f19; padding: 12px; border-radius: 8px; border: 1px solid #334155; }
    .figure-box img { max-width: 100%; border-radius: 6px; box-shadow: 0 4px 12px rgba(0,0,0,0.4); display: block; margin: 0 auto; }
    .figure-caption { font-size: 0.88rem; color: #94a3b8; font-style: italic; margin-top: 8px; font-weight: 600; }
  </style>
</head>
<body>
<div class="container">
  <h1 class="doc-title">MLOps Practical Manual & Command-Line Report</h1>
  <p class="subtitle">Practicals 1 to 4: Model Conversion, Dockerization, AWS EC2 & GCP Cloud Deployment</p>
""")

    current_ss_idx = 0
    for idx, segment in enumerate(segments):
        lines = segment.split('\n')
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
            if re.match(r'^PRACTICAL\s+\d', line_str, re.IGNORECASE):
                html_parts.append(f'<h2 class="practical-title">{line_str}</h2>')
            elif line_str.startswith("AIM:") or line_str.startswith("AIM :"):
                html_parts.append(f'<div class="aim-box"><b>AIM:</b> {line_str[4:].strip()}</div>')
            elif re.match(r'^Step\s+\d', line_str, re.IGNORECASE):
                html_parts.append(f'<h3 class="step-title">{line_str}</h3>')
            elif any(line_str.startswith(c) for c in ["python ", "pip ", "docker ", "sudo ", "curl ", "gcloud ", "aws ", "ssh ", "scp ", "mkdir ", "cd "]):
                html_parts.append(f'<div class="code-box"><code>{line_str}</code></div>')
            else:
                html_parts.append(f'<p>{line_str}</p>')

        if idx < len(segments) - 1 and current_ss_idx < len(SCREENSHOT_FILES):
            ss_filename, ss_caption = SCREENSHOT_FILES[current_ss_idx]
            rel_img = f"screenshots/{ss_filename}"
            html_parts.append(f"""
  <div class="figure-box">
    <img src="{rel_img}" alt="{ss_caption}" />
    <div class="figure-caption">Figure {current_ss_idx + 1}: {ss_caption}</div>
  </div>
""")
            current_ss_idx += 1

    html_parts.append("""
</div>
</body>
</html>
""")

    with open(output_html, 'w', encoding='utf-8') as f:
        f.write("\n".join(html_parts))
    print(f"Successfully generated HTML manual: {output_html}")

if __name__ == "__main__":
    prompt_p = r"C:\Users\Asus\.gemini\antigravity-ide\brain\726a5ff4-339a-4cf9-9118-6f7f3711cbd0\scratch\full_practical_prompt.txt"
    ss_dir = r"c:\Users\Asus\OneDrive\Desktop\Project MLops\screenshots"
    out_docx = r"c:\Users\Asus\OneDrive\Desktop\Project MLops\MLOps_Practicals_Complete_Manual.docx"
    out_html = r"c:\Users\Asus\OneDrive\Desktop\Project MLops\MLOps_Practicals_Complete_Manual.html"

    create_docx_report(prompt_p, ss_dir, out_docx)
    create_html_report(prompt_p, ss_dir, out_html)
